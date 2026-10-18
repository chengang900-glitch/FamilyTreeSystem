"""族谱付印：PDF / Word 三步导出"""
import io
from datetime import datetime, timedelta, timezone

from flask import Blueprint, jsonify, request, send_file

from server.db.conn import query, query_one
from server.service import audit
from server.service.permission import current_user, require_family, require_login
from server.service.printutil import paper_size, register_font

bp = Blueprint('print', __name__, url_prefix='/api/families')


def _load(family_id, root_id, mode, generations):
    fam = query_one('SELECT * FROM family WHERE id=?', (family_id,))
    persons = query('SELECT * FROM person WHERE family_id=? ORDER BY generation, id', (family_id,))
    pmap = {p['id']: dict(p) for p in persons}
    children = {}
    for r in query('SELECT parent_id, child_id FROM parentchild WHERE family_id=?', (family_id,)):
        children.setdefault(r['parent_id'], []).append(r['child_id'])
    spouses = {}
    for m in query('SELECT spouse1_id, spouse2_id FROM marriage WHERE family_id=?', (family_id,)):
        spouses.setdefault(m['spouse1_id'], []).append(m['spouse2_id'])
        spouses.setdefault(m['spouse2_id'], []).append(m['spouse1_id'])

    if mode == 'person' and root_id:
        keep = set()
        stack = [root_id]
        while stack:
            pid = stack.pop()
            if pid in keep or pid not in pmap:
                continue
            keep.add(pid)
            stack.extend(children.get(pid, []))
        items = [pmap[i] for i in sorted(keep, key=lambda x: (pmap[x]['generation'], x))]
    else:
        items = [dict(p) for p in persons]
        if generations and generations > 0:
            base = min((p['generation'] for p in items), default=1)
            items = [p for p in items if p['generation'] < base + generations]
    return fam, items, pmap, children, spouses


def _lines(fam, items, pmap, children, spouses, opts):
    out = []
    by_gen = {}
    for p in items:
        by_gen.setdefault(p['generation'], []).append(p)
    for gen in sorted(by_gen):
        out.append(('h', f"第 {gen} 世"))
        for p in by_gen[gen]:
            name = p['name']
            extra = ' '.join(x for x in (p['courtesy_name'] and f"字{p['courtesy_name']}",
                                         p['art_name'] and f"号{p['art_name']}") if x)
            out.append(('p', f"{name}{('　' + extra) if extra else ''}"))
            born = p['birth_date'] or '生年不详'
            dead = p['death_date'] or '在世'
            out.append(('s', f"{born} — {dead}"))
            if p['bury_address']:
                out.append(('s', f"葬：{p['bury_address']}"))
            if opts.get('include_spouse'):
                for sid in spouses.get(p['id'], []):
                    sp = pmap.get(sid)
                    if sp:
                        out.append(('s', f"配：{sp['name']}"
                                    f"{('（' + sp['birth_date'] + '）') if sp['birth_date'] else ''}"))
            if opts.get('include_children'):
                kids = [pmap[c]['name'] for c in children.get(p['id'], []) if c in pmap]
                if kids:
                    out.append(('s', "子：" + '、'.join(kids)))
            if p['record']:
                out.append(('s', p['record']))
            out.append(('gap', ''))
    return out


# 各类内容行占用的纵向高度（与 print_pdf 的排版保持一致）
LINE_HEIGHTS = {'h': 30, 'p': 16, 's': 13, 'gap': 6}


def _count_pages(content, page_height, margin=42) -> int:
    usable = page_height - 2 * margin - 40
    y, pages = usable, 1
    for kind, _ in content:
        if y < 0:
            pages += 1
            y = usable
        y -= LINE_HEIGHTS.get(kind, 13)
    return pages


@bp.post('/<int:family_id>/print/estimate')
@require_login
@require_family('member')
def estimate(family_id):
    body = request.get_json(silent=True) or {}
    fam, items, pmap, children, spouses = _load(
        family_id, body.get('root_id'), body.get('mode', 'lineage'), body.get('generations'))
    _, height = paper_size(body.get('paper', 'A4'))
    content = _lines(fam, items, pmap, children, spouses, body)
    pages = _count_pages(content, height) + (1 if body.get('add_cover', True) else 0)
    return jsonify({'code': 0, 'msg': 'ok',
                    'data': {'persons': len(items), 'pages': pages}})


@bp.post('/<int:family_id>/print/pdf')
@require_login
@require_family('member')
def print_pdf(family_id):
    from reportlab.lib import colors
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.pdfgen import canvas

    body = request.get_json(silent=True) or {}
    fam, items, pmap, children, spouses = _load(
        family_id, body.get('root_id'), body.get('mode', 'lineage'), body.get('generations'))
    font = register_font()
    width, height = paper_size(body.get('paper', 'A4'))
    margin = 42
    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=(width, height))
    title_font = font
    now = datetime.now(timezone(timedelta(hours=8)))
    page_no = 0

    def new_page():
        nonlocal page_no
        page_no += 1
        c.setFont(font, 9)
        c.setFillColor(colors.HexColor('#888888'))
        c.drawCentredString(width / 2, 24, f"— {page_no} —")
        c.setFillColor(colors.black)
        return margin, height - margin

    x, y = new_page()

    if body.get('add_cover', True):
        cx = width / 2
        title = fam['name'] if fam else '族谱'
        c.setFont(title_font, 26)
        cy = height - height * 0.30
        for i, ch in enumerate(title):
            c.drawCentredString(cx, cy - i * 30, ch)
        c.setFont(font, 11)
        compiler = body.get('compiler') or (fam['compiler'] if fam and fam['compiler'] else '')
        if compiler:
            c.drawCentredString(cx, 150, f"{compiler}　修辑")
        c.drawCentredString(cx, 130, f"{now.strftime('%Y 年 %m 月')}　成谱")
        c.showPage()
        x, y = new_page()

    content = _lines(fam, items, pmap, children, spouses, body)
    for kind, text in content:
        if y < margin + 40:
            c.showPage()
            x, y = new_page()
        if kind == 'h':
            y -= 10
            c.setFont(font, 14)
            c.setFillColor(colors.HexColor('#8c6b4f'))
            c.drawString(margin, y, text)
            c.setFillColor(colors.black)
            y -= 20
        elif kind == 'p':
            c.setFont(font, 12)
            c.drawString(margin, y, text)
            y -= 16
        elif kind == 's':
            c.setFont(font, 9)
            c.setFillColor(colors.HexColor('#555555'))
            c.drawString(margin + 14, y, text)
            c.setFillColor(colors.black)
            y -= 13
        elif kind == 'gap':
            y -= 6
    c.save()
    buf.seek(0)
    audit.log(current_user()['id'], family_id, 'print_pdf', 'family', family_id)
    stamp = now.strftime('%Y%m%d')
    return send_file(buf, as_attachment=True,
                     download_name=f'{(fam["name"] if fam else "族谱")}_{stamp}.pdf',
                     mimetype='application/pdf')


@bp.post('/<int:family_id>/print/docx')
@require_login
@require_family('member')
def print_docx(family_id):
    from docx import Document
    from docx.shared import Pt

    body = request.get_json(silent=True) or {}
    fam, items, pmap, children, spouses = _load(
        family_id, body.get('root_id'), body.get('mode', 'lineage'), body.get('generations'))
    doc = Document()
    style = doc.styles['Normal']
    style.font.name = 'Microsoft YaHei'
    style.font.size = Pt(11)

    now = datetime.now(timezone(timedelta(hours=8)))
    if body.get('add_cover', True):
        doc.add_heading(fam['name'] if fam else '族谱', level=0)
        compiler = body.get('compiler') or (fam['compiler'] if fam and fam['compiler'] else '')
        if compiler:
            doc.add_paragraph(f'{compiler}　修辑')
        doc.add_paragraph(f"{now.strftime('%Y 年 %m 月')}　成谱")
        doc.add_page_break()

    for kind, text in _lines(fam, items, pmap, children, spouses, body):
        if kind == 'h':
            doc.add_heading(text, level=2)
        elif kind == 'p':
            doc.add_paragraph(text)
        elif kind == 's':
            p = doc.add_paragraph(text)
            p.paragraph_format.left_indent = Pt(18)
    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    audit.log(current_user()['id'], family_id, 'print_docx', 'family', family_id)
    return send_file(buf, as_attachment=True,
                     download_name=f'{(fam["name"] if fam else "族谱")}_{now.strftime("%Y%m%d")}.docx',
                     mimetype='application/vnd.openxmlformats-officedocument.wordprocessingml.document')
