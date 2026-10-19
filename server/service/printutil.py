"""付印输出：PDF（ReportLab）与 Word（python-docx）"""
from pathlib import Path

PAPER_MM = {
    'A4': (210, 297),
    'A5': (148, 210),
    '16K': (195, 270),
    '32K': (130, 184),
}

FONT_CANDIDATES = [
    ('CJK', 'C:/Windows/Fonts/msyh.ttc', 0),
    ('CJK', 'C:/Windows/Fonts/simsun.ttc', 0),
    ('CJK', 'C:/Windows/Fonts/simhei.ttf', 0),
    ('CJK', 'C:/Windows/Fonts/simsun.ttf', 0),
    ('CJK', '/usr/share/fonts/opentype/noto/NotoSerifCJK-Regular.ttc', 0),
    ('CJK', '/usr/share/fonts/truetype/noto/NotoSerifCJK-Regular.ttc', 0),
    ('CJK', '/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc', 0),
    ('CJK', '/System/Library/Fonts/Supplemental/Songti.ttc', 0),
    ('CJK', '/System/Library/Fonts/STHeiti Light.ttc', 0),
]

_font_registered = False
_font_name = 'Helvetica'


def register_font():
    """注册中文字体，返回可用字体名"""
    global _font_registered, _font_name
    if _font_registered:
        return _font_name
    try:
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.ttfonts import TTFont
        for name, path, idx in FONT_CANDIDATES:
            if Path(path).exists():
                try:
                    pdfmetrics.registerFont(TTFont(name, path, subfontIndex=idx))
                    _font_name = name
                    break
                except Exception:
                    continue
    except Exception:
        pass
    if _font_name == 'Helvetica':
        # Linux 的 Noto TTC 可能是 ReportLab 不支持的 CFF 轮廓；内置 CJK 字体兜底，
        # 不能静默使用 Helvetica 导致中文变方块。
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.cidfonts import UnicodeCIDFont
        pdfmetrics.registerFont(UnicodeCIDFont('STSong-Light'))
        _font_name = 'STSong-Light'
    _font_registered = True
    return _font_name


def paper_size(key):
    from reportlab.lib.pagesizes import mm
    w, h = PAPER_MM.get((key or 'A4').upper(), PAPER_MM['A4'])
    return w * mm, h * mm


def render_pdf(fam, content, opts):
    """预估与导出共用同一排版器；正文双栏、CJK 折行、长文自动跨栏续页。"""
    if opts.get('template') in ('traditional', 'european'):
        from server.service.print_templates import render_vertical_pdf
        return render_vertical_pdf(fam, content, opts)
    import io
    from datetime import datetime
    from xml.sax.saxutils import escape
    from reportlab.lib.enums import TA_CENTER
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.platypus import (BaseDocTemplate, Frame, PageTemplate, Paragraph,
                                   Spacer, PageBreak, NextPageTemplate)

    font = register_font()
    width, height = paper_size(opts.get('paper', 'A4'))
    margin, gutter = 30, 18
    columns = 1 if opts.get('template') == 'compilation' else 2
    column = (width - margin * 2 - gutter * (columns - 1)) / columns
    buf = io.BytesIO()
    doc = BaseDocTemplate(buf, pagesize=(width, height), title=fam['name'], author=opts.get('compiler') or fam['compiler'] or '')

    def footer(canvas, document):
        canvas.setFont(font, 8)
        canvas.drawCentredString(width / 2, 16, f'— {document.page} —')

    frames = [Frame(margin + i * (column + gutter), margin, column, height - margin * 2,
                    leftPadding=0, rightPadding=0, topPadding=0, bottomPadding=0) for i in range(columns)]
    cover = Frame(margin, margin, width - margin * 2, height - margin * 2,
                  leftPadding=0, rightPadding=0, topPadding=0, bottomPadding=0)
    templates = [PageTemplate(id='body', frames=frames, onPage=footer)]
    if opts.get('add_cover', True):
        templates.insert(0, PageTemplate(id='cover', frames=[cover]))
    doc.addPageTemplates(templates)
    styles = {
        'h': ParagraphStyle('generation', fontName=font, fontSize=13, leading=19,
                            spaceBefore=8, spaceAfter=7, keepWithNext=True, wordWrap='CJK'),
        'p': ParagraphStyle('person', fontName=font, fontSize=11, leading=16,
                            spaceBefore=4, spaceAfter=3, keepWithNext=True, wordWrap='CJK'),
        's': ParagraphStyle('detail', fontName=font, fontSize=9, leading=14,
                            spaceAfter=4, wordWrap='CJK', splitLongWords=True),
    }
    story = []
    if opts.get('add_cover', True):
        title_leading = min(30, max(10, (height - margin * 2 - 130) / max(1, len(fam['name']))))
        title = ParagraphStyle('cover', fontName=font, fontSize=title_leading * .8, leading=title_leading,
                               alignment=TA_CENTER, wordWrap='CJK')
        caption = ParagraphStyle('caption', fontName=font, fontSize=10, leading=16, alignment=TA_CENTER, wordWrap='CJK')
        story += [Spacer(1, 24), Paragraph('<br/>'.join(escape(ch) for ch in fam['name']), title), Spacer(1, 28)]
        compiler = opts.get('compiler') or fam['compiler'] or ''
        if compiler:
            story.append(Paragraph(escape(str(compiler)) + '　修辑', caption))
        story += [Paragraph(datetime.now().strftime('%Y 年 %m 月') + '　成谱', caption),
                  NextPageTemplate('body'), PageBreak()]
    for kind, text in content:
        if kind == 'gap':
            story.append(Spacer(1, 6))
        else:
            story.append(Paragraph(escape(str(text)).replace('\n', '<br/>'), styles[kind]))
    if not content:
        story.append(Paragraph('暂无成员资料', styles['s']))
    doc.build(story)
    return buf.getvalue(), doc.page


def render_docx(fam, content, opts):
    import io
    from datetime import datetime
    from docx import Document
    from docx.shared import Mm, Pt
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    doc = Document()
    section = doc.sections[0]
    w, h = PAPER_MM[opts.get('paper', 'A4')]
    section.page_width, section.page_height = Mm(w), Mm(h)
    section.top_margin = section.bottom_margin = Mm(15)
    section.left_margin = section.right_margin = Mm(12)
    normal = doc.styles['Normal']
    normal.font.name = 'Microsoft YaHei'
    normal.font.size = Pt(10)
    normal.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
    normal.paragraph_format.space_after = Pt(4)
    if opts.get('add_cover', True):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Mm(18)
        run = p.add_run('\n'.join(fam['name']))
        run.font.size = Pt(24)
        compiler = opts.get('compiler') or fam['compiler'] or ''
        if compiler:
            doc.add_paragraph(str(compiler) + '　修辑').alignment = WD_ALIGN_PARAGRAPH.CENTER
        doc.add_paragraph(datetime.now().strftime('%Y 年 %m 月') + '　成谱').alignment = WD_ALIGN_PARAGRAPH.CENTER
        doc.add_page_break()

    table, details = None, None
    for kind, text in content:
        if opts.get('template') == 'compilation':
            if kind == 'h':
                doc.add_heading(str(text), level=2)
            elif kind == 'p':
                doc.add_heading(str(text), level=3)
            elif kind == 's':
                doc.add_paragraph(str(text))
            continue
        if kind == 'h':
            doc.add_heading(text, level=2)
            table = doc.add_table(rows=1, cols=2)
            table.style = 'Table Grid'
            table.autofit = False
            table.columns[0].width = Mm((w - 24) * .25)
            table.columns[1].width = Mm((w - 24) * .75)
            table.rows[0].cells[0].text, table.rows[0].cells[1].text = '族人', '生平与亲属'
            repeat = OxmlElement('w:tblHeader')
            table.rows[0]._tr.get_or_add_trPr().append(repeat)
        elif kind == 'p' and table is not None:
            name, details = table.add_row().cells
            name.text = str(text)
            details.text = ''
        elif kind == 's' and details is not None:
            if details.text:
                details.add_paragraph(str(text))
            else:
                details.text = str(text)
    if not content:
        doc.add_paragraph('暂无成员资料')
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    field = OxmlElement('w:fldSimple')
    field.set(qn('w:instr'), 'PAGE')
    footer._p.append(field)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
