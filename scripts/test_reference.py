"""参考功能专项回归：独立临时库、真实认证、实际导出文件。"""
import io
import json
import sys
import zipfile
from pathlib import Path

from _testenv import bootstrap, assert_isolated, relax_bootstrap_admin, cleanup

bootstrap('reference')
from server.app import app
from server.db.conn import execute, query, query_one
from server.service import settings
from server.blueprint.print_bp import _load
from flask import g
from docx import Document
from pypdf import PdfReader

assert_isolated(app)
relax_bootstrap_admin(app)
PASS, FAIL = [], []


def check(name, value, detail=''):
    (PASS if value else FAIL).append(name)
    print(('PASS ' if value else 'FAIL ') + name + (f' {detail}' if detail else ''))


c = app.test_client()
c.post('/api/auth/login', json={'username': 'admin', 'password': 'admin123'})
c.put('/api/admin/settings', json={'security.rate_limit_per_minute': 10000})
fid = c.post('/api/admin/families', json={'name': '参考源谱', 'generation_chars': '天地玄黄',
    'generation_frozen': True, 'compiler': '测试修辑者'}).get_json()['data']['id']
base = f'/api/families/{fid}'
ids = []
for gen, name, kind in [(1, '始祖', 1), (1, '配偶', 3), (2, '儿子', 1), (2, '女儿', 2), (3, '孙子', 1)]:
    ids.append(c.post(base + '/persons', json={'generation': gen, 'name': name, 'type': kind,
        'birth_place': '杭州', 'courtesy_name': '测试字', 'record': '生平长文。' * 500}).get_json()['data']['id'])
dad, mom, son, daughter, grandson = ids
c.post(base + '/marriages', json={'spouse1_id': dad, 'spouse2_id': mom,
    'marriage_date': '1950-01-01', 'divorce_date': '1960-01-01', 'marriage_place': '杭州'})
c.post(base + '/parentchild', json={'parent_id': dad, 'child_ids': [son, daughter]})
c.post(base + '/parentchild', json={'parent_id': son, 'child_ids': [grandson]})
resource = c.post(base + '/resources', data={'name': '旧谱', 'type': '1',
    'file': (io.BytesIO(b'sample'), 'sample.txt')}).get_json()['data']['id']
c.post(base + f'/resources/{resource}/relations', json={'entity_type': 'person', 'entity_id': dad})

csv = zipfile.ZipFile(io.BytesIO(c.get(base + '/export?format=csv').data))
check('CSV 九模块', len(csv.namelist()) == 9 and 'resource_relations.csv' in csv.namelist())
data = json.loads(c.get(base + '/export?format=json').data)
for import_meta in [False, True]:
    target = c.post('/api/admin/families', json={'name': f'目标{import_meta}'}).get_json()['data']['id']
    r = c.post(f'/api/families/{target}/import', data={'file': (io.BytesIO(json.dumps(data).encode()), 'x.json'),
        'include_metadata': str(import_meta).lower()})
    f = c.get(f'/api/families/{target}').get_json()['data']
    check(f'元数据导入开关 {import_meta}', r.status_code == 200 and
          (f['generation_chars'] == '天地玄黄' if import_meta else f['name'] == '目标False'))

ged_zip = zipfile.ZipFile(io.BytesIO(c.get(base + '/export?format=gedcom').data))
ged = ged_zip.read(ged_zip.namelist()[0])
check('GEDCOM 婚姻日期地点', b'1 MARR' in ged and b'1950-01-01' in ged and '杭州'.encode() in ged)
target = c.post('/api/admin/families', json={'name': 'GED目标'}).get_json()['data']['id']
r = c.post(f'/api/families/{target}/import', data={'file': (io.BytesIO(ged), 'x.ged')})
with app.app_context():
    marr = dict(query_one('SELECT * FROM marriage WHERE family_id=?', (target,)))
    pp = dict(query_one('SELECT * FROM person WHERE family_id=? ORDER BY id', (target,)))
check('GEDCOM 婚姻往返', marr['marriage_date'] == '1950-01-01' and marr['divorce_date'] == '1960-01-01'
      and marr['marriage_place'] == '杭州')
check('GEDCOM 人物往返', pp['courtesy_name'] == '测试字' and pp['birth_place'] == '杭州' and pp['record'] == '生平长文。' * 500)

for paper, dims in [('A4', (210, 297)), ('A5', (148, 210)), ('16K', (195, 270)), ('32K', (130, 184))]:
    r = c.post(base + '/print/docx', json={'paper': paper, 'include_children': True})
    doc = Document(io.BytesIO(r.data))
    sec = doc.sections[-1]
    check(f'Word 纸张 {paper}', abs(sec.page_width.mm - dims[0]) < .1 and abs(sec.page_height.mm - dims[1]) < .1)
    check(f'Word 分栏及子女区分 {paper}', bool(doc.tables) and '女：女儿' in '\n'.join(
        cell.text for table in doc.tables for row in table.rows for cell in row.cells))

with app.test_request_context():
    g.current_user = {'id': 1, 'is_superuser': True}
    check('个人模式世代深度', [p['id'] for p in _load(fid, dad, 'person', 1)[1]] == [dad])

opts = {'paper': '32K', 'include_children': True, 'include_spouse': True, 'add_cover': True}
r = c.post(base + '/print/pdf', json=opts)
check('长文 PDF 可生成', r.status_code == 200 and r.data.startswith(b'%PDF'))
if r.status_code == 200:
    pdf = PdfReader(io.BytesIO(r.data))
    estimate = c.post(base + '/print/estimate', json=opts).get_json()['data']
    check('预估页数等于实际 PDF', estimate['pages'] == len(pdf.pages), str((estimate['pages'], len(pdf.pages))))
    out = Path('/tmp/fts-improve-20261019')
    out.mkdir(exist_ok=True)
    (out / 'sample-32K.pdf').write_bytes(r.data)
    (out / 'sample-32K.docx').write_bytes(c.post(base + '/print/docx', json=opts).data)

r = c.post(base + '/parentchild', json={'father_id': dad, 'mother_id': mom,
    'children': [{'child_id': son, 'father_type': 1, 'mother_type': 3},
                 {'child_id': daughter, 'father_type': 5, 'mother_type': 6}]})
check('父母同屏逐子女登记', r.status_code == 200)
with app.app_context():
    types = {(p['parent_id'], p['child_id']): p['type'] for p in query('SELECT * FROM parentchild WHERE family_id=?', (fid,))}
check('逐子女生养关系', types.get((mom, son)) == 3 and types.get((mom, daughter)) == 6)
with app.app_context():
    before = query_one('SELECT COUNT(*) c FROM parentchild WHERE family_id=?', (fid,))['c']
r = c.post(base + '/parentchild', json={'father_id': grandson, 'children': [{'child_id': dad, 'father_type': 1}]})
check('批量亲子拒绝环路', r.status_code == 400)
with app.app_context():
    check('批量失败不落半条数据', query_one('SELECT COUNT(*) c FROM parentchild WHERE family_id=?', (fid,))['c'] == before)

for entity_type in ['person', 'marriage', 'parentchild', 'event', 'branch']:
    r = c.get(base + f'/relation-targets?entity_type={entity_type}&keyword=始祖&page=1&size=20')
    check(f'文献对象远程搜索 {entity_type}', r.status_code == 200 and 'items' in (r.get_json().get('data') or {}))

# 无效批量输入必须整批回滚，不得覆盖其他家谱。
other = c.post('/api/admin/families', json={'name': '隔离谱'}).get_json()['data']['id']
foreign = c.post(f'/api/families/{other}/persons', json={'name': '外谱成员', 'generation': 1}).get_json()['data']['id']
r = c.post(base + '/parentchild', json={'mother_id': mom, 'children': [
    {'child_id': grandson}, {'child_id': foreign}]})
check('批量亲子拒绝跨谱', r.status_code in (400, 403))
r = c.post(base + '/parentchild', json={'father_id': dad, 'mother_id': dad, 'children': [{'child_id': son}]})
check('批量亲子拒绝同一父母', r.status_code == 400)
r = c.post(base + '/parentchild', json={'father_id': dad, 'children': [{'child_id': son}, {'child_id': son}]})
check('批量亲子拒绝重复子女', r.status_code == 400)

# 文献各类对象都必须能搜索到第一页以后的记录。
from server.common import now_str
ts = now_str()
with app.app_context():
    for i in range(205):
        person = execute('INSERT INTO person(family_id,generation,name,type,created_at,updated_at) VALUES (?,?,?,?,?,?)',
                         (fid, 2, f'后续族人{i:03}', 1, ts, ts))
        execute('INSERT INTO marriage(family_id,spouse1_id,spouse2_id,marriage_order,created_at,updated_at) VALUES (?,?,?,?,?,?)',
                (fid, mom, person, i + 1, ts, ts))
        execute('INSERT INTO parentchild(family_id,parent_id,child_id,created_at,updated_at) VALUES (?,?,?,?,?)',
                (fid, dad, person, ts, ts))
        execute('INSERT INTO branch(family_id,name,created_at,updated_at) VALUES (?,?,?,?)', (fid, f'后续分支{i:03}', ts, ts))
        execute('INSERT INTO event(family_id,name,time,created_at,updated_at) VALUES (?,?,?,?,?)',
                (fid, f'后续事件{i:03}', '2000', ts, ts))
for kind, keyword in [('person','后续族人204'), ('marriage','后续族人204'), ('parentchild','后续族人204'),
                      ('branch','后续分支204'), ('event','后续事件204')]:
    r = c.get(base + '/relation-targets', query_string={'entity_type': kind, 'keyword': keyword, 'size': 20})
    check(f'文献可搜索第201条之后 {kind}', r.get_json()['data']['total'] == 1)
check('文献对象类型白名单', c.get(base + '/relation-targets?entity_type=users').status_code == 400)

# 元数据非法时，原有业务记录必须保留；权限、状态、slug 不随文件导入。
bad = {'family': {'name': '替换失败', 'generation_chars': {}}, 'persons': []}
r = c.post(base + '/import', data={'file': (io.BytesIO(json.dumps(bad).encode()), 'bad.json'),
                                  'strategy': 'overwrite', 'include_metadata': 'true'})
check('非法项目资料被拒绝', r.status_code == 400)
check('非法覆盖导入保留原数据', c.get(base + '/persons?size=1').get_json()['data']['total'] == 210)
pack = io.BytesIO()
with zipfile.ZipFile(pack, 'w') as z:
    z.writestr('data.json', json.dumps({'family': {'name': '可导入资料', 'slug': '恶意目录', 'status': 0,
                                                'generation_chars': '测试字辈', 'generation_frozen': 1}}))
r = c.post(f'/api/families/{other}/import', data={'file': (io.BytesIO(pack.getvalue()), 'data.zip'), 'include_metadata': 'true'})
meta = c.get(f'/api/families/{other}').get_json()['data']
check('ZIP 可导入项目资料', r.status_code == 200 and meta['generation_chars'] == '测试字辈')
check('项目导入不更改权限和目录', meta['status'] == 1 and meta['slug'] != '恶意目录')

# 逐亲子类型和隐私扩展往返，第三方仅识别标准标签时仍有基本关系。
z = zipfile.ZipFile(io.BytesIO(c.get(base + '/export?format=gedcom').data))
kindged = z.read(z.namelist()[0])
kindtarget = c.post('/api/admin/families', json={'name': '类型往返'}).get_json()['data']['id']
c.post(f'/api/families/{kindtarget}/import', data={'file': (io.BytesIO(kindged), 'types.ged')})
with app.app_context():
    pc = query('SELECT a.name parent,b.name child,pc.type FROM parentchild pc JOIN person a ON pc.parent_id=a.id '
               'JOIN person b ON pc.child_id=b.id WHERE pc.family_id=?', (kindtarget,))
check('GEDCOM 逐父母类型往返', any(r['parent']=='配偶' and r['child']=='女儿' and r['type']==6 for r in pc))
for opts in [{'mode': 'person'}, {'mode': 'person', 'root_id': foreign}, {'compiler': {}}, {'paper': 'B0'}]:
    check(f'付印拒绝无效参数 {opts}', c.post(base + '/print/docx', json=opts).status_code == 400)

print(f'结果：通过 {len(PASS)}，失败 {len(FAIL)}')
cleanup()
sys.exit(bool(FAIL))
