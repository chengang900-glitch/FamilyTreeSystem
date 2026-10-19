"""付印模板隔离验收；需测试依赖 pypdf，不连接生产库。"""
import io
from pathlib import Path
from docx import Document
from pypdf import PdfReader
from _testenv import bootstrap, assert_isolated, relax_bootstrap_admin, cleanup

bootstrap('print_templates')
from server.app import app
from server.db.conn import query_one
assert_isolated(app)
relax_bootstrap_admin(app)
passed = 0


def check(label, condition):
    global passed
    assert condition, label
    passed += 1
    print('PASS', label, flush=True)


def pdf_text(blob):
    pdf = PdfReader(io.BytesIO(blob))
    return pdf, ''.join(''.join(page.extract_text().split()) for page in pdf.pages)


c = app.test_client()
assert c.post('/api/auth/login', json={'username': 'admin', 'password': 'admin123'}).status_code == 200
c.put('/api/admin/settings', json={'security.rate_limit_per_minute': 10000, 'storage.backup_enabled': False})
fid = c.post('/api/admin/families', json={'name': '林氏宗谱', 'hall_name': '敦睦堂', 'compiler': '验收修辑者'}).json['data']['id']
base = f'/api/families/{fid}'
ids = []
names = ['林继金', '周玉兰', '林世信', '林与君', '林兴胜', '林立权', '林延庆']
for i, name in enumerate(names):
    r = c.post(base + '/persons', json={'name': name, 'generation': 16 if i < 2 else 15 + i,
        'type': 3 if i == 1 else 1, 'record': '勤俭持家，生平依原始资料记载。', 'birth_date': '1900-01-01'})
    assert r.status_code == 200, r.json
    ids.append(r.json['data']['id'])
assert c.post(base + '/marriages', json={'spouse1_id': ids[0], 'spouse2_id': ids[1], 'marriage_order': 1}).status_code == 200
assert c.post(base + '/parentchild', json={'father_id': ids[0], 'mother_id': ids[1], 'children': [{'child_id': ids[2]}]}).status_code == 200
out = Path('output/print-templates-20261019');out.mkdir(parents=True, exist_ok=True)
for template in ['original', 'traditional', 'european', 'compilation']:
    for paper, dims in [('A4', (210, 297)), ('A5', (148, 210)), ('16K', (195, 270)), ('32K', (130, 184))]:
        opts = {'template': template, 'paper': paper, 'include_spouse': True, 'include_children': True, 'add_cover': True}
        response = c.post(base + '/print/pdf', json=opts)
        assert response.status_code == 200, response.json
        pdf, text = pdf_text(response.data)
        estimate = c.post(base + '/print/estimate', json=opts).json['data']
        check(f'{template}/{paper} 纸张、实际页数与人数一致', len(pdf.pages) == estimate['pages'] and estimate['persons'] == 7
              and abs(float(pdf.pages[0].mediabox.width) * 25.4 / 72 - dims[0]) < .1
              and abs(float(pdf.pages[0].mediabox.height) * 25.4 / 72 - dims[1]) < .1)
        check(f'{template}/{paper} 姓名与生平完整', all(name in text for name in names) and text.count('勤俭持家') >= 7)
        if template != 'original':
            check(f'{template}/{paper} 显示真实双亲', '父母：林继金、周玉兰' in text)
        if paper == 'A4':
            (out / (template + '-sample.pdf')).write_bytes(response.data)
    opts = {'template': template, 'add_cover': False}
    no_cover = c.post(base + '/print/pdf', json=opts)
    pdf, text = pdf_text(no_cover.data)
    check(f'{template} 不加封面时页数一致', len(pdf.pages) == c.post(base + '/print/estimate', json=opts).json['data']['pages'])

# 长文每个字符只来自一个真实字段，检查续页时没有截断或复制正文。
long_record = '卷首独有字琼' + '璇' * 1100 + '卷尾独有字珏'
c.put(base + f'/persons/{ids[0]}', json={'record': long_record})
for template in ['traditional', 'european', 'compilation']:
    opts = {'template': template, 'paper': '32K', 'add_cover': False}
    response = c.post(base + '/print/pdf', json=opts)
    assert response.status_code == 200, response.json
    pdf, text = pdf_text(response.data)
    check(f'{template} 长文跨页不丢字不复制', len(pdf.pages) > 1 and text.count('璇') == 1100 and text.count('琼') == 1 and text.count('珏') == 1)
    check(f'{template} 长文页数准确', len(pdf.pages) == c.post(base + '/print/estimate', json=opts).json['data']['pages'])
    (out / (template + '-long-32K.pdf')).write_bytes(response.data)

for template in ['original', 'compilation']:
    r = c.post(base + '/print/docx', json={'template': template, 'add_cover': False})
    assert r.status_code == 200, r.json
    doc = Document(io.BytesIO(r.data))
    text = '\n'.join(p.text for p in doc.paragraphs) if template == 'compilation' else '\n'.join(cell.text for t in doc.tables for row in t.rows for cell in row.cells)
    check(f'{template} Word结构与内容完整', bool(doc.tables) == (template == 'original') and all(name in text for name in names) and long_record in text)
    (out / (template + '-sample.docx')).write_bytes(r.data)
for template in ['traditional', 'european']:
    check(f'{template} 不支持Word时明确拒绝', c.post(base + '/print/docx', json={'template': template}).status_code == 400)
for value in ['unknown', None, {}, [], 12]:
    check(f'拒绝非法模板 {value!r}', c.post(base + '/print/pdf', json={'template': value}).status_code == 400)
check('旧请求未传模板仍使用原版Word', bool(Document(io.BytesIO(c.post(base + '/print/docx', json={}).data)).tables))

empty = c.post('/api/admin/families', json={'name': '空谱'}).json['data']['id']
for template in ['original', 'traditional', 'european', 'compilation']:
    r = c.post(f'/api/families/{empty}/print/pdf', json={'template': template, 'add_cover': False})
    check(f'{template} 空谱可导出', r.status_code == 200 and '暂无成员资料' in pdf_text(r.data)[1])

uid_r = c.post('/api/admin/users', json={'username': 'print_reader', 'password': 'test1234', 'must_change_pwd': False})
assert uid_r.status_code == 200, uid_r.json
with app.app_context():
    uid = query_one('SELECT id FROM users WHERE username=?', ('print_reader',))['id']
c.post(base + '/members', json={'user_id': uid, 'role': 'member'})
c.put(base + f'/persons/{ids[0]}', json={'is_private': True, 'record': 'PRIVATE_RECORD_TOKEN', 'phone': 'PRIVATE_PHONE_TOKEN'})
member = app.test_client();member.post('/api/auth/login', json={'username': 'print_reader', 'password': 'test1234'})
anonymous = app.test_client()
for template in ['original', 'traditional', 'european', 'compilation']:
    r = member.post(base + '/print/pdf', json={'template': template, 'include_spouse': True})
    text = pdf_text(r.data)[1] if r.status_code == 200 else ''
    check(f'{template} 普通成员脱敏且禁止跨谱', r.status_code == 200 and 'PRIVATE_RECORD_TOKEN' not in text and 'PRIVATE_PHONE_TOKEN' not in text
          and member.post(f'/api/families/{empty}/print/pdf', json={'template': template}).status_code == 403)
    check(f'{template} 未登录不可导出', anonymous.post(base + '/print/pdf', json={'template': template}).status_code == 401)
opts = {'template': 'traditional', 'mode': 'person', 'root_id': ids[2], 'generations': 1, 'add_cover': False}
r = c.post(base + '/print/pdf', json=opts)
pdf, text = pdf_text(r.data)
check('个人模式世代深度保持生效', c.post(base + '/print/estimate', json=opts).json['data']['persons'] == 1 and names[2] in text and names[3] not in text)
print(f'付印模板检查：{passed} 项通过', flush=True)
cleanup()
